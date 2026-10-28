import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None, color="CBD5E1", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="{top or 'single'}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{left or 'single'}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{bottom or 'single'}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{right or 'single'}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

def add_header_footer(doc):
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    
    # Header
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("TRƯỜNG ĐẠI HỌC CNTT & TT THÁI NGUYÊN (ICTU) | ĐỒ ÁN TỐT NGHIỆP 2026")
    hrun.font.name = "Arial"
    hrun.font.size = Pt(8.5)
    hrun.font.color.rgb = RGBColor(100, 116, 139)
    
    # Footer
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = fp.add_run("Hệ Thống Đặt Sân Bóng Tự Động Tích Hợp AI & Ví Tiền Virtual Wallet — Nhóm 02")
    frun.font.name = "Arial"
    frun.font.size = Pt(8.5)
    frun.font.color.rgb = RGBColor(100, 116, 139)

def style_table(table, col_widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(table.rows):
        is_header = (i == 0)
        bg_color = "1E3A8A" if is_header else ("F8FAFC" if i % 2 == 1 else "FFFFFF")
        
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if is_header:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            
        for j, cell in enumerate(row.cells):
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=130, right=130)
            set_cell_borders(cell, color="CBD5E1", sz="4")
            
            if col_widths and j < len(col_widths):
                cell.width = Inches(col_widths[j])
                
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                for run in p.runs:
                    run.font.name = "Arial"
                    if is_header:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.size = Pt(9)
                    else:
                        run.font.size = Pt(9)
                        run.font.color.rgb = RGBColor(30, 41, 59)

def add_code_block(doc, code_text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.6)
    
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
    set_cell_borders(cell, left="single", top="none", right="none", bottom="none", color="0284C7", sz="24")
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    
    run = p.add_run(code_text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(3)

def add_markdown_paragraph(p, text, is_bold=False, is_italic=False, default_color=RGBColor(51, 65, 85)):
    # Parse markdown inline formatting: **bold**, *italic*, `code`
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = p.add_run(token[2:-2])
            run.font.name = "Arial"
            run.font.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(15, 23, 42)
        elif token.startswith('*') and token.endswith('*'):
            run = p.add_run(token[1:-1])
            run.font.name = "Arial"
            run.font.italic = True
            run.font.size = Pt(10)
            run.font.color.rgb = default_color
        elif token.startswith('`') and token.endswith('`'):
            run = p.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(194, 65, 12)
        else:
            run = p.add_run(token)
            run.font.name = "Arial"
            run.font.bold = is_bold
            run.font.italic = is_italic
            run.font.size = Pt(10)
            run.font.color.rgb = default_color

def build_docx_from_markdown(md_path, docx_path):
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    add_header_footer(doc)
    
    # ---------------------------------------------------------
    # TRANG BÌA (COVER PAGE)
    # ---------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(36)
    title_p.paragraph_format.space_after = Pt(4)
    
    r_inst = title_p.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN & TRUYỀN THÔNG THÁI NGUYÊN\n")
    r_inst.font.name = "Arial"
    r_inst.font.size = Pt(13)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(30, 58, 138)
    
    r_sub = title_p.add_run("KHOA CÔNG NGHỆ THÔNG TIN — ĐỒ ÁN TỐT NGHIỆP ĐẠI HỌC\n\n")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(100, 116, 139)
    
    cover_box = doc.add_table(rows=1, cols=1)
    cover_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_cell = cover_box.cell(0, 0)
    c_cell.width = Inches(6.5)
    set_cell_background(c_cell, "F1F5F9")
    set_cell_margins(c_cell, top=240, bottom=240, left=240, right=240)
    set_cell_borders(c_cell, top="single", bottom="single", left="single", right="single", color="1E3A8A", sz="16")
    
    cp = c_cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    r_rp = cp.add_run("BÁO CÁO THỰC THI & PHÁT TRIỂN ỨNG DỤNG CHUYÊN SÂU\n\n")
    r_rp.font.name = "Arial"
    r_rp.font.size = Pt(12)
    r_rp.font.bold = True
    r_rp.font.color.rgb = RGBColor(2, 132, 199)
    
    r_main_title = cp.add_run("HỆ THỐNG ĐẶT SÂN BÓNG TỰ ĐỘNG\nTÍCH HỢP AI & VÍ ĐIỆN TỬ VIRTUAL WALLET\n\n")
    r_main_title.font.name = "Arial"
    r_main_title.font.size = Pt(18)
    r_main_title.font.bold = True
    r_main_title.font.color.rgb = RGBColor(30, 58, 138)
    
    r_desc = cp.add_run("Kiến trúc Microservices (.NET 8/10 Web API + Python AI + SQL Server)\nThuật toán Dynamic Pricing, Trợ lý ảo LM Studio & Thanh toán VietQR\n")
    r_desc.font.name = "Arial"
    r_desc.font.size = Pt(10)
    r_desc.font.italic = True
    r_desc.font.color.rgb = RGBColor(71, 85, 105)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(24)
    
    info_table = doc.add_table(rows=4, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_data = [
        ("Sinh viên thực hiện:", "1. Nguyễn Văn A (Trưởng nhóm - KT Phần mềm)\n2. Nguyễn Thị B (Thành viên - HT Thông tin)"),
        ("Lớp / Khóa:", "K20 - Kỹ thuật Phần mềm / Công nghệ Thông tin"),
        ("Giảng viên hướng dẫn:", "TS. Giảng Viên Hướng Dẫn"),
        ("Thời gian thực hiện:", "27/07/2026 – 27/09/2026 (9 Tuần)")
    ]
    for idx, (label, val) in enumerate(info_data):
        c0 = info_table.cell(idx, 0)
        c1 = info_table.cell(idx, 1)
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(3)
        p0.paragraph_format.space_after = Pt(3)
        r0 = p0.add_run(label)
        r0.font.name = "Arial"
        r0.font.bold = True
        r0.font.size = Pt(10)
        r0.font.color.rgb = RGBColor(30, 58, 138)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(3)
        p1.paragraph_format.space_after = Pt(3)
        r1 = p1.add_run(val)
        r1.font.name = "Arial"
        r1.font.size = Pt(10)
        r1.font.color.rgb = RGBColor(30, 41, 59)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(40)
    
    ft_p = doc.add_paragraph()
    ft_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_city = ft_p.add_run("THÁI NGUYÊN – NĂM 2026")
    r_city.font.name = "Arial"
    r_city.font.size = Pt(11)
    r_city.font.bold = True
    r_city.font.color.rgb = RGBColor(100, 116, 139)
    
    doc.add_page_break()
    
    # ---------------------------------------------------------
    # DYNAMIC TABLE OF CONTENTS (MỤC LỤC TỰ ĐỘNG WORD)
    # ---------------------------------------------------------
    toc_head = doc.add_paragraph()
    toc_head.paragraph_format.space_before = Pt(16)
    toc_head.paragraph_format.space_after = Pt(6)
    r_th = toc_head.add_run("MỤC LỤC BÁO CÁO ĐỒ ÁN TỐT NGHIỆP")
    r_th.font.name = "Arial"
    r_th.font.size = Pt(15)
    r_th.font.bold = True
    r_th.font.color.rgb = RGBColor(30, 58, 138)
    
    pPr = toc_head._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="4" w:color="0284C7"/></w:pBdr>')
    pPr.append(pBdr)
    
    note_p = doc.add_paragraph()
    note_p.paragraph_format.space_before = Pt(4)
    note_p.paragraph_format.space_after = Pt(12)
    r_note = note_p.add_run("💡 Hướng dẫn: Trong Microsoft Word, nhấp chuột phải vào vùng mục lục bên dưới và chọn 'Update Field' (hoặc nhấn phím F9) để cập nhật chính xác số trang theo bản in.")
    r_note.font.name = "Arial"
    r_note.font.size = Pt(9)
    r_note.font.italic = True
    r_note.font.color.rgb = RGBColor(2, 132, 199)
    
    # Word Dynamic TOC field
    p_toc = doc.add_paragraph()
    p_toc.paragraph_format.space_after = Pt(18)
    run_toc = p_toc.add_run()
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> TOC \o "1-3" \h \z \u </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run_toc._r.append(fldChar1)
    run_toc._r.append(instrText)
    run_toc._r.append(fldChar2)
    run_toc._r.append(fldChar3)
    
    doc.add_page_break()
    
    # ---------------------------------------------------------
    # PARSE MARKDOWN CONTENT & RENDER DOCUMENT
    # ---------------------------------------------------------
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    in_code_block = False
    code_buffer = []
    code_lang = ""
    
    in_table = False
    table_lines = []
    
    for idx, line in enumerate(lines):
        raw_line = line.rstrip('\r\n')
        stripped = raw_line.strip()
        
        # Check code block toggle
        if stripped.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_lang = stripped[3:].strip()
                code_buffer = []
                continue
            else:
                in_code_block = False
                add_code_block(doc, "\n".join(code_buffer))
                code_buffer = []
                continue
                
        if in_code_block:
            code_buffer.append(raw_line)
            continue
            
        # Check table
        if stripped.startswith('|') and stripped.endswith('|'):
            if not in_table:
                in_table = True
                table_lines = [stripped]
            else:
                table_lines.append(stripped)
            continue
        else:
            if in_table:
                # Render table
                in_table = False
                # Filter out separator line |---|---|
                data_rows = [r for r in table_lines if not re.match(r'^\|(\s*:?-+:?\s*\|)+$', r)]
                if data_rows:
                    parsed_rows = []
                    for r in data_rows:
                        cells = [c.strip() for c in r.strip('|').split('|')]
                        parsed_rows.append(cells)
                    
                    num_cols = max(len(row) for row in parsed_rows)
                    t = doc.add_table(rows=len(parsed_rows), cols=num_cols)
                    
                    for r_i, row in enumerate(parsed_rows):
                        for c_i, val in enumerate(row):
                            if c_i < num_cols:
                                cell = t.cell(r_i, c_i)
                                p = cell.paragraphs[0]
                                p.paragraph_format.space_before = Pt(2)
                                p.paragraph_format.space_after = Pt(2)
                                p.paragraph_format.line_spacing = 1.15
                                # replace <br> with newline
                                clean_val = val.replace('<br>', '\n').replace('<br/>', '\n')
                                add_markdown_paragraph(p, clean_val, is_bold=(r_i == 0))
                                
                    # Auto column widths
                    total_width = 6.6
                    col_w = [total_width / num_cols] * num_cols
                    # Custom heuristics for specific tables
                    if num_cols == 3:
                        col_w = [1.2, 3.8, 1.6]
                    elif num_cols == 4:
                        col_w = [0.8, 2.4, 2.2, 1.2]
                    elif num_cols == 5:
                        col_w = [0.8, 1.6, 2.2, 1.1, 0.9]
                    elif num_cols == 7:
                        col_w = [0.7, 1.0, 1.2, 1.1, 1.0, 1.0, 0.6]
                        
                    style_table(t, col_widths=col_w)
                    sp = doc.add_paragraph()
                    sp.paragraph_format.space_before = Pt(0)
                    sp.paragraph_format.space_after = Pt(4)
                table_lines = []
                
        # Empty line
        if not stripped:
            continue
            
        # Horizontal rule
        if stripped == '---':
            continue
            
        # Headings
        if stripped.startswith('# '):
            text = stripped[2:].strip()
            # If it's a major Part (PHẦN), add page break
            if "PHẦN" in text:
                doc.add_page_break()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = "Arial"
            run.font.bold = True
            run.font.size = Pt(15)
            run.font.color.rgb = RGBColor(30, 58, 138)
            pPr = p._p.get_or_add_pPr()
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="4" w:color="0284C7"/></w:pBdr>')
            pPr.append(pBdr)
            continue
            
        if stripped.startswith('## '):
            text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(13)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = "Arial"
            run.font.bold = True
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(2, 132, 199)
            continue
            
        if stripped.startswith('### '):
            text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = "Arial"
            run.font.bold = True
            run.font.size = Pt(11.5)
            run.font.color.rgb = RGBColor(3, 105, 161)
            continue
            
        if stripped.startswith('#### '):
            text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = "Arial"
            run.font.bold = True
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(15, 23, 42)
            continue
            
        # Blockquote >
        if stripped.startswith('> '):
            text = stripped[2:].strip()
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.cell(0, 0)
            cell.width = Inches(6.6)
            set_cell_background(cell, "EFF6FF")
            set_cell_margins(cell, top=80, bottom=80, left=140, right=140)
            set_cell_borders(cell, left="single", top="none", right="none", bottom="none", color="3B82F6", sz="16")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            add_markdown_paragraph(p, text, is_italic=True, default_color=RGBColor(30, 58, 138))
            sp = doc.add_paragraph()
            sp.paragraph_format.space_before = Pt(0)
            sp.paragraph_format.space_after = Pt(3)
            continue
            
        # Bullet list
        if stripped.startswith('* ') or stripped.startswith('- '):
            bullet_text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2.5)
            p.paragraph_format.line_spacing = 1.15
            add_markdown_paragraph(p, bullet_text)
            continue
            
        # Numbered list item like 1. 2.
        m_num = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if m_num:
            num_prefix = m_num.group(1) + ". "
            num_text = m_num.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.left_indent = Inches(0.25)
            r_num = p.add_run(num_prefix)
            r_num.font.name = "Arial"
            r_num.font.bold = True
            r_num.font.size = Pt(10)
            r_num.font.color.rgb = RGBColor(30, 58, 138)
            add_markdown_paragraph(p, num_text)
            continue
            
        # Regular paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        add_markdown_paragraph(p, stripped)
        
    doc.save(docx_path)
    print(f"Successfully generated DOCX report at: {docx_path}")

if __name__ == '__main__':
    md_file = r"C:\Users\klqdo\Downloads\1websanbongai\BAO_CAO_CHI_TIET_DO_AN_TOT_NGHIEP.md"
    docx_file = r"C:\Users\klqdo\Downloads\1websanbongai\BAO_CAO_DO_AN_TOT_NGHIEP_SAN_BONG_AI.docx"
    build_docx_from_markdown(md_file, docx_file)
