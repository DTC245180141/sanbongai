import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document

docx_path = r"c:\Users\klqdo\Downloads\BAO_CAO_DO_AN_TOT_NGHIEP_SAN_BONG_AI_FULL_PLANTUML (1).docx"
doc = Document(docx_path)

print('Total paragraphs:', len(doc.paragraphs))
print('Total tables:', len(doc.tables))
print()
print('=== ALL PARAGRAPHS (first 100 non-empty) ===')
count = 0
for i, para in enumerate(doc.paragraphs):
    if para.text.strip():
        print(f'[{i}] [{para.style.name}] {para.text[:120]}')
        count += 1
        if count >= 100:
            break
